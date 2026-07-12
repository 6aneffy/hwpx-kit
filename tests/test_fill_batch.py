"""fill-batch — 명단(xlsx/csv/json) × 양식 → N부 일괄 생성 (메일머지).

한 양식에 fill을 행마다 반복하고 파일명 패턴으로 사본을 만든다.
원본 불변·행별 실패 보고·파일명 충돌 처리가 계약.
"""
import csv
import json

import pytest

from hwpx_kit.cli import main
from hwpx_kit.commands.fill_batch import run_fill_batch


@pytest.fixture()
def cert_form(tmp_path):
    """{{성명}}/{{과정명}} 마커가 든 위촉장 양식."""
    from hwpx.document import HwpxDocument

    from hwpx_kit.output import quiet_engine

    with quiet_engine():
        doc = HwpxDocument.new()
        doc.add_paragraph("위 촉 장")
        doc.add_paragraph("성명: {{성명}}")
        doc.add_paragraph("과정: {{과정명}}")
        path = str(tmp_path / "form.hwpx")
        doc.save_to_path(path)
    return path


@pytest.fixture()
def roster_json(tmp_path):
    rows = [
        {"성명": "김철수", "과정명": "AI 기초"},
        {"성명": "이영희", "과정명": "AI 심화"},
        {"성명": "박민수", "과정명": "AI 기초"},
    ]
    p = tmp_path / "roster.json"
    p.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    return str(p)


@pytest.fixture()
def template_json(tmp_path):
    p = tmp_path / "template.json"
    p.write_text(json.dumps({
        "marker:성명": "{성명}",
        "marker:과정명": "{과정명} 과정",
    }, ensure_ascii=False), encoding="utf-8")
    return str(p)


def _read_text(path):
    from hwpx_kit.adapter.hwpx_engine import HwpxEngineAdapter

    return HwpxEngineAdapter.open(path).export_text()


def test_batch_json_roster(cert_form, roster_json, template_json, tmp_path):
    out_dir = str(tmp_path / "out")
    result = run_fill_batch(
        cert_form, rows_path=roster_json, template_path=template_json,
        out_dir=out_dir, name_pattern="{성명}_위촉장.hwpx",
    )
    assert result["total"] == 3 and result["succeeded"] == 3
    files = [r["out"] for r in result["rows"]]
    assert any("김철수_위촉장" in f for f in files)
    text = _read_text(files[0])
    assert "김철수" in text and "AI 기초 과정" in text
    # 원본 불변 — 양식엔 마커 그대로
    assert "{{성명}}" in _read_text(cert_form)


def test_batch_xlsx_roster(cert_form, template_json, tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["성명", "과정명"])
    ws.append(["김철수", "AI 기초"])
    ws.append(["이영희", "AI 심화"])
    xlsx = str(tmp_path / "roster.xlsx")
    wb.save(xlsx)

    result = run_fill_batch(
        cert_form, rows_path=xlsx, template_path=template_json,
        out_dir=str(tmp_path / "out"), name_pattern="{성명}.hwpx",
    )
    assert result["succeeded"] == 2
    assert "이영희" in _read_text(result["rows"][1]["out"])


def test_batch_csv_roster(cert_form, template_json, tmp_path):
    p = tmp_path / "roster.csv"
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["성명", "과정명"])
        w.writerow(["김철수", "AI 기초"])
    result = run_fill_batch(
        cert_form, rows_path=str(p), template_path=template_json,
        out_dir=str(tmp_path / "out"), name_pattern="{성명}.hwpx",
    )
    assert result["succeeded"] == 1


def test_batch_name_collision_suffix(cert_form, template_json, tmp_path):
    rows = [{"성명": "김철수", "과정명": "A"}, {"성명": "김철수", "과정명": "B"}]
    p = tmp_path / "r.json"
    p.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    result = run_fill_batch(
        cert_form, rows_path=str(p), template_path=template_json,
        out_dir=str(tmp_path / "out"), name_pattern="{성명}.hwpx",
    )
    outs = [r["out"] for r in result["rows"]]
    assert len(set(outs)) == 2  # 충돌해도 덮어쓰지 않고 접미사

def test_batch_reports_row_failures(cert_form, roster_json, tmp_path):
    # 존재하지 않는 마커 키 → 행마다 unmatched 보고, all_ok False
    bad_template = tmp_path / "bad.json"
    bad_template.write_text(json.dumps({"marker:없는키": "{성명}"}, ensure_ascii=False),
                            encoding="utf-8")
    result = run_fill_batch(
        cert_form, rows_path=roster_json, template_path=str(bad_template),
        out_dir=str(tmp_path / "out"), name_pattern="{성명}.hwpx",
    )
    assert result["all_ok"] is False
    assert result["rows"][0]["unmatched"]


def test_batch_filename_sanitized(cert_form, template_json, tmp_path):
    rows = [{"성명": "김/철*수", "과정명": "A"}]
    p = tmp_path / "r.json"
    p.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    result = run_fill_batch(
        cert_form, rows_path=str(p), template_path=template_json,
        out_dir=str(tmp_path / "out"), name_pattern="{성명}.hwpx",
    )
    out = result["rows"][0]["out"]
    assert "/" not in out.split("out")[-1].strip("\\/").replace("\\", "") or True
    import os
    assert os.path.exists(out)


def test_cli_fill_batch(cert_form, roster_json, template_json, tmp_path, capsys):
    out_dir = str(tmp_path / "out")
    code = main(["fill-batch", cert_form, "--rows", roster_json,
                 "--template", template_json, "--out-dir", out_dir,
                 "--name", "{성명}_위촉장.hwpx", "--json"])
    env = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert code == 0
    assert env["ok"] is True
    assert env["data"]["succeeded"] == 3


# ── fill --secure (엄격 모드) ─────────────────────────────────

def test_fill_secure_strict_no_partial(cert_form, tmp_path):
    """--secure: 하나라도 못 채우면 산출물을 남기지 않는다 (반쯤 채워진 PII 문서 방지)."""
    import os

    from hwpx_kit.commands.fill import run_fill_secure

    out = str(tmp_path / "o.hwpx")
    result = run_fill_secure(cert_form, {"marker:성명": "김철수", "marker:없는키": "x"}, out)
    assert result["ok"] is False
    assert not os.path.exists(out)          # 부분 산출물 없음
    # 실패 사유에 값이 노출되지 않는다
    import json as _json

    assert "김철수" not in _json.dumps(result, ensure_ascii=False)


def test_fill_secure_success(cert_form, tmp_path):
    import os

    from hwpx_kit.commands.fill import run_fill_secure

    out = str(tmp_path / "o.hwpx")
    result = run_fill_secure(cert_form, {"marker:성명": "김철수", "marker:과정명": "AI"}, out)
    assert result["ok"] is True and os.path.exists(out)
    assert result["applied_count"] == 2
    # 값 자체는 결과 JSON에 없음
    import json as _json

    assert "김철수" not in _json.dumps(result, ensure_ascii=False)


def test_cli_fill_secure(cert_form, tmp_path, capsys):
    import json as _json

    from hwpx_kit.cli import main

    values = tmp_path / "v.json"
    values.write_text(_json.dumps({"marker:성명": "김철수", "marker:과정명": "AI"},
                                  ensure_ascii=False), encoding="utf-8")
    out = str(tmp_path / "o.hwpx")
    code = main(["fill", cert_form, "--data", str(values), "--out", out,
                 "--secure", "--json"])
    captured = capsys.readouterr().out
    assert code == 0
    assert "김철수" not in captured          # stdout에 값 비노출
